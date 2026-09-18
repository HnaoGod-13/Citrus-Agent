"""Editable Word export for project reports."""
from __future__ import annotations

from pathlib import Path
import re
from collections.abc import Iterable
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


CHINESE_FONT = "宋体"
LATIN_FONT = "Times New Roman"
BODY_SIZE = 12
HEADING_SIZE = 16
TABLE_SIZE = 10.5


def _typeface(r_pr, size: float) -> None:
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    # Theme fonts can override the explicit fonts in uploaded Word templates.
    fonts.attrib.clear()
    for script in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{script}"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CHINESE_FONT)
    for tag in ("sz", "szCs"):
        node = r_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            r_pr.append(node)
        node.set(qn("w:val"), str(round(size * 2)))


def _flush_left(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ind = paragraph._p.get_or_add_pPr().get_or_add_ind()
    ind.attrib.clear()
    for attr in ("left", "right", "firstLine", "leftChars", "rightChars", "firstLineChars"):
        ind.set(qn(f"w:{attr}"), "0")


def _set_defaults(document: Document, *, preserve_page_setup: bool = False) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _typeface(normal.element.get_or_add_rPr(), BODY_SIZE)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 22)] + [(f"Heading {level}", HEADING_SIZE) for level in range(1, 10)]:
        style = styles[name]
        _typeface(style.element.get_or_add_rPr(), size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.left_indent = Pt(0)
        style.paragraph_format.right_indent = Pt(0)
        if name != "Title":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style.element.get_or_add_pPr().remove(borders)
    if not preserve_page_setup:
        for section in document.sections:
            section.top_margin = Cm(2.4)
            section.bottom_margin = Cm(2.2)
            section.left_margin = Cm(2.8)
            section.right_margin = Cm(2.4)


def _paragraphs(document: Document) -> Iterable:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _replace_placeholders(document: Document, values: dict[str, str]) -> None:
    for paragraph in _paragraphs(document):
        original = paragraph.text
        updated = original
        for key, value in values.items():
            updated = updated.replace("{{" + key + "}}", value)
        if updated != original:
            # Word frequently splits one placeholder across several runs.  Set
            # the complete paragraph so those placeholders are still filled.
            paragraph.text = updated


def _section_values(markdown: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {}
    current = ""
    for raw in str(markdown or "").splitlines():
        line = raw.strip()
        if line.startswith("## "):
            current = line[3:].strip()
            sections.setdefault(current, [])
            continue
        if current:
            sections[current].append(raw)
    values: dict[str, str] = {}
    for name, lines in sections.items():
        output = []
        for raw in lines:
            line = raw.strip()
            if not line or re.match(r"^\|?\s*:?-{3,}", line):
                continue
            if line.startswith("|"):
                line = "｜".join(cell.strip() for cell in line.strip("|").split("|"))
            line = re.sub(r"^###\s+", "", line)
            line = re.sub(r"^[-*]\s+", "• ", line)
            line = re.sub(r"^\d+[.)]\s+", "", line)
            line = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
            line = re.sub(r"\[([^]]+)\]\((https?://[^)]+)\)", r"\1（\2）", line)
            output.append(line)
        values[name] = "\n".join(output)
    return values


def _without_template_sections(markdown: str, used_sections: set[str]) -> str:
    output: list[str] = []
    include = False
    for raw in str(markdown or "").splitlines():
        line = raw.strip()
        if line.startswith("## "):
            include = line[3:].strip() not in used_sections
        elif line.startswith("# "):
            include = False
        if include:
            output.append(raw)
    return "\n".join(output).strip()


def _field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    for node in (begin, code, separate, text, end):
        run.append(node)
    paragraph._p.append(run)


def _add_page_number(section, *, preserve_existing: bool = False) -> None:
    footer = section.footer.paragraphs[0]
    if preserve_existing and footer.text.strip():
        return
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = "柑橘产业项目报告  第 "
    _field(footer, " PAGE ")
    footer.add_run(" 页")


def _add_toc(document: Document, headings: list[str]) -> None:
    document.add_heading("目录", level=1)
    for index, heading in enumerate(headings, 1):
        paragraph = document.add_paragraph(f"{index:02d}  {heading}")
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.runs[0].font.size = Pt(BODY_SIZE)
    document.add_page_break()


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement("w:t"); node.text = text
    run.append(node); hyperlink.append(run); paragraph._p.append(hyperlink)


def _fill_with_links(paragraph, text: str) -> None:
    cursor = 0
    for match in re.finditer(r"https?://[^\s]+", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        url = match.group(0).rstrip("。，；,.;)")
        _add_hyperlink(paragraph, url, url)
        cursor = match.start() + len(url)
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _table_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}"); margins.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def _add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Normal Table"
    table.autofit = True
    for row_index, values in enumerate(rows):
        row_properties = table.rows[row_index]._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index].strip() if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _table_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _borders(properties, tag: str, visible: dict[str, int]) -> None:
    for old in properties.findall(qn(f"w:{tag}")):
        properties.remove(old)
    borders = OxmlElement(f"w:{tag}")
    for edge in ("top", "bottom", "left", "right", "start", "end", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single" if edge in visible else "nil")
        if edge in visible:
            border.set(qn("w:sz"), str(visible[edge]))
            border.set(qn("w:color"), "000000")
            border.set(qn("w:space"), "0")
        borders.append(border)
    properties.append(borders)


def _three_line_table(table) -> None:
    properties = table.tblPr
    for tag in ("tblStyle", "tblLook"):
        for old in properties.findall(qn(f"w:{tag}")):
            properties.remove(old)
    # Clear template banding and direct fills as well as the old grid borders.
    for node in table.xpath(".//w:shd | .//w:highlight | .//w:cnfStyle | .//w:pBdr"):
        node.getparent().remove(node)
    _borders(properties, "tblBorders", {"top": 12, "bottom": 12})
    rows = table.findall(qn("w:tr"))
    for index, row in enumerate(rows):
        row_pr = row.get_or_add_trPr()
        if row_pr.find(qn("w:cantSplit")) is None:
            row_pr.append(OxmlElement("w:cantSplit"))
        # Do not mark the first row as a repeating header.  The report format
        # intentionally lets a table continue on the next page without
        # printing the same heading row again.
        for repeat_header in row_pr.findall(qn("w:tblHeader")):
            row_pr.remove(repeat_header)
        visible = {"top": 12, "bottom": 6} if index == 0 else {}
        if index == len(rows) - 1:
            visible["bottom"] = 12
        for cell in row.findall(qn("w:tc")):
            _borders(cell.get_or_add_tcPr(), "tcBorders", visible)
            if index == 0:
                # Keep the first data row with the table heading when a table
                # begins close to a page boundary.  This avoids an orphaned
                # heading while still keeping the heading non-repeating.
                for paragraph in cell.xpath(".//w:p"):
                    p_pr = paragraph.get_or_add_pPr()
                    p_pr.get_or_add_keepNext().val = True


def _paragraph_kind(paragraph: Paragraph) -> str:
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        if style.name == "Title":
            return "title"
        if style.name.startswith("Heading "):
            return "heading"
        style = style.base_style
    outline = paragraph._p.xpath("./w:pPr/w:outlineLvl/@w:val")
    return "heading" if outline and outline[0] in {str(i) for i in range(9)} else "body"


def _apply_report_format(document: Document) -> None:
    """Enforce the report format on generated text and uploaded templates alike."""
    parts = {document.part.partname: document.element}
    for section in document.sections:
        for name in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
            story = getattr(section, name)
            if not story.is_linked_to_previous:
                parts[story.part.partname] = story.part.element
    for root in parts.values():
        for table in root.xpath(".//w:tbl"):
            _three_line_table(table)
        for element in root.xpath(".//w:p"):
            paragraph = Paragraph(element, document)
            in_table = bool(element.xpath("ancestor::w:tc"))
            kind = _paragraph_kind(paragraph)
            size = TABLE_SIZE if in_table else HEADING_SIZE if kind == "heading" else 22 if kind == "title" else BODY_SIZE
            if in_table or kind == "heading":
                _flush_left(paragraph)
            if kind == "heading" and not in_table:
                paragraph.paragraph_format.keep_with_next = True
                for text in element.xpath(".//w:t"):
                    if text.text:
                        text.text = text.text.lstrip(" \t\u3000")
                        if text.text:
                            break
            # Include hyperlink runs and field results (page numbers), which
            # paragraph.runs does not always expose.
            p_pr = element.get_or_add_pPr()
            mark = p_pr.find(qn("w:rPr"))
            if mark is None:
                mark = OxmlElement("w:rPr")
                p_pr.append(mark)
            _typeface(mark, size)
            for run in element.xpath(".//w:r"):
                r_pr = run.get_or_add_rPr()
                _typeface(r_pr, size)
                if in_table or kind in {"title", "heading"}:
                    color = r_pr.find(qn("w:color"))
                    if color is None:
                        color = OxmlElement("w:color")
                        r_pr.append(color)
                    color.attrib.clear()
                    color.set(qn("w:val"), "000000")


def markdown_to_docx(markdown: str, output_path: str | Path, *, profile: dict[str, Any] | None = None, template_path: str | Path | None = None, sources: list[dict[str, Any]] | None = None) -> Path:
    profile = profile or {}
    template = Path(template_path) if template_path else None
    document = Document(str(template)) if template and template.is_file() and template.suffix.lower() == ".docx" else Document()
    using_template = bool(template and template.is_file() and template.suffix.lower() == ".docx")
    _set_defaults(document, preserve_page_setup=using_template)
    lines = str(markdown or "").splitlines()
    contents = [line[3:].strip() for line in lines if line.strip().startswith("## ")]
    if using_template:
        section_values = _section_values(markdown)
        template_text = "\n".join(paragraph.text for paragraph in _paragraphs(document))
        used_sections = {name for name in section_values if "{{" + name + "}}" in template_text}
        aliases = {
            **{k: str(v or "") for k, v in profile.items()},
            **section_values,
            "项目名称": str(profile.get("title") or "柑橘产业项目报告"),
            "项目标题": str(profile.get("title") or "柑橘产业项目报告"),
            "使用单位": str(profile.get("agency") or ""),
            "承办部门": str(profile.get("department") or ""),
            "经办人员": str(profile.get("preparedBy") or ""),
            "统计区域": str(profile.get("region") or ""),
            "报告周期": str(profile.get("period") or ""),
            "工作目的": str(profile.get("purpose") or ""),
        }
        _replace_placeholders(document, aliases)
        if used_sections:
            markdown = _without_template_sections(markdown, used_sections)
            lines = markdown.splitlines()
        if any(line.strip() for line in lines):
            document.add_section(WD_SECTION.NEW_PAGE)
    else:
        cover = document.add_paragraph(style="Title")
        cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover.paragraph_format.first_line_indent = Cm(0)
        cover.add_run(str(profile.get("title") or "柑橘产业项目报告")).bold = True
        meta = document.add_paragraph(
            "\n".join(filter(None, [str(profile.get("agency") or ""), str(profile.get("department") or ""), str(profile.get("period") or "")]))
        )
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.paragraph_format.first_line_indent = Cm(0)
        meta.paragraph_format.space_before = Pt(120)
        document.add_page_break()
        _add_toc(document, contents)

    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        index += 1
        if not line:
            continue
        if line.startswith("|") and index < len(lines) and re.match(r"^\s*\|?\s*:?-{3,}", lines[index]):
            rows = [[cell.strip() for cell in line.strip("|").split("|")]]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            _add_markdown_table(document, rows)
            continue
        if line.startswith("# "):
            if line[2:].strip() != str(profile.get("title") or "").strip():
                document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^[-*]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            _fill_with_links(paragraph, re.sub(r"^[-*]\s+", "", line))
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif re.match(r"^\d+[.)]\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            _fill_with_links(paragraph, re.sub(r"^\d+[.)]\s+", "", line))
            paragraph.paragraph_format.first_line_indent = Cm(0)
        else:
            paragraph = document.add_paragraph()
            _fill_with_links(paragraph, line)
    if sources and not re.search(r"^##\s+参考资料\s*$", str(markdown or ""), flags=re.MULTILINE):
        document.add_heading("来源索引（可点击）", level=2)
        for index, source in enumerate(sources, 1):
            title = str(source.get("title") or source.get("document_title") or "未命名来源")
            url = str(source.get("url") or source.get("source") or "")
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(f"{title} · ")
            if url.startswith(("http://", "https://")):
                _add_hyperlink(paragraph, url, url)
            else:
                paragraph.add_run(url or "本地证据")
    for section in document.sections:
        _add_page_number(section, preserve_existing=using_template)
    _apply_report_format(document)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    return out


__all__ = ["markdown_to_docx"]
