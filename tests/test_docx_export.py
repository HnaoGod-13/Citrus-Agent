from docx import Document
from docx.shared import Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.reporting.docx_export import markdown_to_docx


def test_custom_word_template_preserves_layout_and_fills_section_placeholders(tmp_path):
    template_path = tmp_path / "unit-template.docx"
    template = Document()
    template.sections[0].top_margin = Cm(3)
    template.sections[0].header.paragraphs[0].text = "单位项目申报材料"
    title = template.add_paragraph()
    title.add_run("{{项目")
    title.add_run("名称}}")
    template.add_heading("摘要", level=1)
    template.add_paragraph("{{项目摘要}}")
    template.save(template_path)

    markdown = """# 武鸣沃柑项目

## 项目摘要
这是由批次事实和外部证据形成的摘要。

## 项目背景
这是需要追加到模板的项目背景。
"""
    output_path = tmp_path / "report.docx"
    markdown_to_docx(
        markdown,
        output_path,
        profile={"title": "武鸣沃柑项目"},
        template_path=template_path,
    )

    output = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "武鸣沃柑项目" in text
    assert "这是由批次事实和外部证据形成的摘要。" in text
    assert "这是需要追加到模板的项目背景。" in text
    assert "{{" not in text
    assert output.sections[0].header.paragraphs[0].text == "单位项目申报材料"
    assert round(output.sections[0].top_margin.cm, 1) == 3.0


def assert_typeface(element, half_points):
    runs = element.xpath(".//w:r")
    assert runs
    for run in runs:
        fonts = run.find(qn("w:rPr")).find(qn("w:rFonts"))
        assert fonts.get(qn("w:eastAsia")) == "宋体"
        assert fonts.get(qn("w:ascii")) == "Times New Roman"
        assert fonts.get(qn("w:hAnsi")) == "Times New Roman"
        assert not any("Theme" in key for key in fonts.attrib)
        assert run.xpath("./w:rPr/w:sz/@w:val") == [str(half_points)]


def assert_three_line_table(table):
    assert not table._tbl.xpath(".//w:shd | .//w:highlight")
    visible = table._tbl.xpath("./w:tblPr/w:tblBorders/*[@w:val='single']")
    assert {node.tag for node in visible} == {qn("w:top"), qn("w:bottom")}
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            visible = cell._tc.xpath("./w:tcPr/w:tcBorders/*[@w:val='single']")
            expected = {qn("w:top"), qn("w:bottom")} if row_index == 0 else {qn("w:bottom")} if row_index == len(table.rows)-1 else set()
            assert {node.tag for node in visible} == expected
            assert_typeface(cell._tc, 21)
            for paragraph in cell.paragraphs:
                assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
                assert paragraph.paragraph_format.first_line_indent == 0
            assert cell._tc.xpath(".//w:rPr/w:color/@w:val") == ["000000"] * len(cell._tc.xpath(".//w:r"))


def test_generated_report_uses_requested_fonts_sizes_and_three_line_tables(tmp_path):
    output_path = tmp_path / "formatted.docx"
    markdown_to_docx(
        "# 示例报告\n\n## 项目摘要\n中文 NFC 12.8 °Brix 正文。\n"
        "\n### 建设建议\n第二段正文 https://example.org/reference\n"
        "\n| 结论层级 | 当前结论 | 使用边界 |\n|---|---|---|\n"
        "| 已有批次事实 | 沃柑 20 吨 | 检测复核 |\n"
        "| 系统建议 | NFC 果汁 | 需要小试 |\n"
        "| 投资决策 | 待测算 | 待询价 |\n",
        output_path, profile={"title": "示例报告"},
    )
    doc = Document(output_path)
    for paragraph in doc.paragraphs:
        if not paragraph.text:
            continue
        if paragraph.style.name.startswith("Heading"):
            assert_typeface(paragraph._p, 32)
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT
            assert paragraph.paragraph_format.first_line_indent == 0
        elif paragraph.style.name != "Title":
            assert_typeface(paragraph._p, 24)
    toc = next(p for p in doc.paragraphs if p.text.startswith("01  "))
    assert_typeface(toc._p, 24)
    assert_three_line_table(doc.tables[0])
    assert_typeface(doc.sections[0].footer._element, 24)
    assert doc.sections[0].footer._element.xpath(".//w:instrText[contains(text(), 'PAGE')]")
    assert len([r for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink')]) == 1


def test_uploaded_template_direct_formatting_cannot_override_report_format(tmp_path):
    template_path = tmp_path / "colored-template.docx"
    template = Document()
    heading = template.add_heading("　　项目摘要", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(1)
    fonts = heading.runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:asciiTheme"), "majorHAnsi")
    fonts.set(qn("w:eastAsiaTheme"), "majorEastAsia")
    template.add_paragraph("{{项目摘要}}")
    template.sections[0].header.paragraphs[0].text = "报告 Header 2026"
    template.sections[0].footer.paragraphs[0].text = "页脚 Footer"
    template.sections[0].different_first_page_header_footer = True
    template.sections[0].first_page_header.paragraphs[0].text = "首页 Header"
    template.sections[0].even_page_footer.paragraphs[0].text = "偶数页 Footer"
    table = template.add_table(rows=3, cols=2)
    table.style = "Light Shading Accent 1"
    for row in table.rows:
        for cell in row.cells:
            cell.text = "中文 NFC 20"
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "1F4E78")
            cell._tc.get_or_add_tcPr().append(shading)
    template.save(template_path)
    output_path = tmp_path / "template-result.docx"
    markdown_to_docx("## 项目摘要\n模板正文。", output_path, template_path=template_path)
    doc = Document(output_path)
    heading = doc.paragraphs[0]
    assert heading.text == "项目摘要"
    assert_typeface(heading._p, 32)
    assert heading.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert heading.paragraph_format.first_line_indent == 0
    for name in ("header", "footer", "first_page_header", "even_page_footer"):
        assert_typeface(getattr(doc.sections[0], name)._element, 24)
    assert_three_line_table(doc.tables[0])
